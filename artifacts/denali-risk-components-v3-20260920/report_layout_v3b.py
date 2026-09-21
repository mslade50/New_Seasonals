"""Layout for the v3b rebuild.

Cosmetics copied verbatim from artifacts/denali-risk-introduction-20260917 so the
rebuilt pages look like the ones the team already has: black headings, F1F1F1
header shading, the tighter cell spacing. Two changes only: the footer date, and
the page break now rides on the first paragraph of each page (a standalone break
paragraph can be pushed off a full page and leave a blank sheet behind it).
"""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT.parents[1]
stats=json.loads((SOURCE/'data/signal_horizon_stats.json').read_text())['signals']
FOOTER='DENALI  |  Internal  |  21 September 2026'

def para(d,text,style=None):
    p=d.add_paragraph(style=style)
    for i,t in enumerate(re.split(r'\*\*(.*?)\*\*',text,flags=re.S)):
        r=p.add_run(t);r.bold=bool(i%2)
    return p

def build(name,title,pages,compact=False):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.48 if compact else .62);s.bottom_margin=Inches(.48 if compact else .60)
    s.left_margin=s.right_margin=Inches(.60)
    s.footer_distance=Inches(.22)
    for x in list(d.styles.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
    for nm,sz in [('Normal',10 if compact else 10.5),('Title',21),('Heading 1',13)]:
        st=d.styles[nm];st.font.name='Calibri';st.font.size=Pt(sz)
        st.font.color.rgb=RGBColor.from_string('000000')
        st.paragraph_format.line_spacing=1.02 if compact else 1.07
        st.paragraph_format.space_after=Pt(5 if compact else 7)
        if nm=='Heading 1':st.paragraph_format.space_before=Pt(8)
    s.footer.paragraphs[0].text=FOOTER
    s.footer.paragraphs[0].runs[0].font.size=Pt(8)
    if not compact:
        s.footer.paragraphs[0].add_run('  |  ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');s.footer.paragraphs[0]._p.append(fld)
    d.core_properties.title=title;d.core_properties.author='Denali'
    d.add_paragraph(title,'Title');md=['# '+title,'']
    for n,page in enumerate(pages):
        pending=bool(n)
        if n:md.append('\n---\n')
        for b in page:
            lead=None
            if b[0]=='h':lead=d.add_heading(b[1],1);md.extend(['## '+b[1],''])
            elif b[0] in ['p','small']:
                lead=para(d,b[1]);md.extend([b[1],''])
                if b[0]=='small':
                    for r in lead.runs:r.font.size=Pt(8.5 if compact else 9)
            elif b[0]=='table':
                if pending:
                    lead=d.add_paragraph()
                    for attr,val in [('space_before',Pt(0)),('space_after',Pt(0)),('line_spacing',1)]:
                        setattr(lead.paragraph_format,attr,val)
                    lead.add_run().font.size=Pt(1)
                headers,rows,widths=b[1:];t=d.add_table(rows=1,cols=len(headers));t.autofit=False
                for c,w in zip(t.columns,widths):c.width=Inches(w)
                for row,values in [(t.rows[0],headers)]+[(t.add_row(),v) for v in rows]:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for ci,(cell,text,w) in enumerate(zip(row.cells,values,widths)):
                        cell.width=Inches(w);p=cell.paragraphs[0]
                        for j,part in enumerate(re.split(r'\*\*(.*?)\*\*',str(text),flags=re.S)):
                            run=p.add_run(part);run.bold=bool(j%2) or row is t.rows[0]
                            run.font.size=Pt(9.5 if compact else 10)
                        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4 if compact else 5)
                        p.paragraph_format.line_spacing=1.01 if compact else 1.05
                        tc=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                        for side in ['left','right']:
                            e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
                for cell in t.rows[0].cells:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F1F1');cell._tc.get_or_add_tcPr().append(sh)
                    for r in cell.paragraphs[0].runs:r.bold=True
                md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
                md.extend('| '+' | '.join(str(x).replace('\n',' / ') for x in r)+' |' for r in rows);md.append('')
            if pending and lead is not None:
                lead.paragraph_format.page_break_before=True;pending=False
    d.save(ROOT/(name+'.docx'));(ROOT/(name+'.md')).write_text('\n'.join(md),encoding='utf-8')
    print(name, 'words:',len(' '.join(md).split()))
